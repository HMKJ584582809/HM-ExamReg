# -*- coding: utf-8 -*-
"""修正「导出示例-计算机类考试-报名导入模板.xlsx」里遗留的占位/非法值。

原示例数据中有 3 行残留了占位文本或字典外的值，导致照示例填写反而导入失败：
  * 报考科目 = '报考科目'（占位文本，不在科目字典内）
  * 考试机构编码 = '机构编码' / 考点编码 = '机构编码08'（占位文本，格式非法）
本脚本把这些单元格修正为字典内/格式合法的值，其余示例数据原样保留。

用法：
  python tools/fix_template_values.py <输出目录>
"""
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.config import load_dicts                              # noqa: E402
from app.routers.applications import COMPUTER_FIELDS           # noqa: E402
from app.routers.export import build_workbook                  # noqa: E402
from openpyxl import load_workbook                             # noqa: E402

FNAME = "导出示例-计算机类考试-报名导入模板.xlsx"
SITE_RE = re.compile(r"^[A-Za-z0-9]{4,12}$")
ORG_RE = re.compile(r"^[A-Za-z0-9]{2,20}$")


def read_rows(src: Path, fields: list):
    if not src.exists():
        return []
    ws = load_workbook(src).worksheets[0]
    it = ws.iter_rows(values_only=True)
    next(it, None)
    rows = []
    for row in it:
        item = {}
        for i, f in enumerate(fields):
            v = row[i] if i < len(row) else None
            item[f] = "" if v is None else str(v)
        if any(item.values()):
            rows.append(item)
    return rows


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    target = out / FNAME
    d = load_dicts()
    subjects = d["computer"]["subjects"]
    rows = read_rows(target, COMPUTER_FIELDS)
    fixed_subject = fixed_site = fixed_org = 0

    for i, r in enumerate(rows):
        if r.get("subject") not in subjects:
            r["subject"] = subjects[i % len(subjects)]
            fixed_subject += 1
        org = r.get("org_code", "")
        if not ORG_RE.match(org or ""):
            org = "3201"
            r["org_code"] = org
            fixed_org += 1
        if not SITE_RE.match(r.get("exam_site_code", "") or ""):
            r["exam_site_code"] = (org + "01")[:12]
            fixed_site += 1

    exam = {"id": 0, "exam_type": "computer", "name": target.stem}
    cnt = build_workbook(exam, rows, target, len(rows))
    print("已修正：%s（数据行 %d；科目 %d 处、机构编码 %d 处、考点编码 %d 处）"
          % (target.name, cnt, fixed_subject, fixed_org, fixed_site))
    return 0


if __name__ == "__main__":
    sys.exit(main())
