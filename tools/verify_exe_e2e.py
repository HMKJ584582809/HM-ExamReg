# -*- coding: utf-8 -*-
"""
exe 端到端验证：创建考试 → 批量导入 → 报名列表 → 汇总导出
------------------------------------------------------------------
在冻结（exe）环境里跑通「管理权限级别的批量导入」完整链路，重点验证：
  * openpyxl 在打包环境下可正常读写 xlsx
  * python-multipart 在打包环境下可正常接收文件上传
  * 导入后数据进入报名列表，并可汇总导出

用法：
  python tools/verify_exe_e2e.py --base http://127.0.0.1:8801
"""
import argparse
import io
import json
import sys
import urllib.error
import urllib.request
import uuid

PASS = 0
FAIL = 0
BASE = "http://127.0.0.1:8801"
TOKEN = ""


def call(method, path, body=None, raw=False, token=None):
    data = json.dumps(body, ensure_ascii=False).encode() if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    tk = token if token is not None else TOKEN
    if tk:
        req.add_header("Authorization", "Bearer " + tk)
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            b = r.read()
            return (r.status, b) if raw else (r.status, json.loads(b.decode("utf-8", "replace")))
    except urllib.error.HTTPError as e:
        b = e.read()
        if raw:
            return e.code, b
        try:
            return e.code, json.loads(b.decode("utf-8", "replace"))
        except Exception:
            return e.code, {"raw": b[:300].decode("utf-8", "replace")}


def upload(path, fields, filename, content):
    boundary = "----wb" + uuid.uuid4().hex
    buf = io.BytesIO()
    for k, v in fields.items():
        buf.write(("--%s\r\nContent-Disposition: form-data; name=\"%s\"\r\n\r\n%s\r\n"
                   % (boundary, k, v)).encode("utf-8"))
    buf.write(("--%s\r\nContent-Disposition: form-data; name=\"file\";"
               " filename=\"%s\"\r\nContent-Type: application/vnd.openxmlformats-"
               "officedocument.spreadsheetml.sheet\r\n\r\n" % (boundary, filename)).encode("utf-8"))
    buf.write(content)
    buf.write(("\r\n--%s--\r\n" % boundary).encode("utf-8"))
    req = urllib.request.Request(BASE + path, data=buf.getvalue(), method="POST")
    req.add_header("Content-Type", "multipart/form-data; boundary=" + boundary)
    if TOKEN:
        req.add_header("Authorization", "Bearer " + TOKEN)
    try:
        with urllib.request.urlopen(req, timeout=300) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        b = e.read()
        try:
            return e.code, json.loads(b.decode("utf-8"))
        except Exception:
            return e.code, {"raw": b[:300].decode("utf-8", "ignore")}


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  [OK]   " + name)
    else:
        FAIL += 1
        print("  [FAIL] " + name + ("  -> " + str(extra) if extra else ""))


def main():
    global BASE, TOKEN
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://127.0.0.1:8801")
    args = ap.parse_args()
    BASE = args.base

    from openpyxl import Workbook, load_workbook

    print("=" * 66)
    print("exe 端到端验证 · " + BASE)
    print("=" * 66)

    # 1. 登录
    print("\n[1] 登录与清理历史临时批次")
    st, lg = call("POST", "/api/auth/login",
                  {"account": "admin", "password": "Admin@123", "remember": True})
    check("admin 登录成功", lg.get("code") == 0, lg.get("message"))
    TOKEN = (lg.get("data") or {}).get("access_token") or ""
    check("取得令牌", bool(TOKEN))

    # 清理上一轮遗留的临时批次，保证幂等
    def _leftovers():
        st, ex = call("GET", "/api/exams?page=1&page_size=100")
        return [it for it in ((ex.get("data") or {}).get("list") or [])
                if str(it.get("name", "")).startswith("E2E-")]

    for it in _leftovers():
        call("DELETE", "/api/exams/%d" % it["id"])
    if _leftovers():
        # 非草稿批次不可删除，用系统清理兜底（仅用于验证环境的空库）
        call("POST", "/api/system/reset",
             {"scopes": ["exams", "applications"], "confirm": "确认清空"})

    # 2. 创建考试批次
    print("\n[2] 创建考试批次（计算机类 / 普通话）")
    import time
    tag = time.strftime("%H%M%S")
    st, ex1 = call("POST", "/api/exams", {
        "exam_type": "computer", "exam_year": 2026, "exam_month": 9,
        "name": "E2E-计算机-%s" % tag, "signup_start_at": "2026-09-01 00:00",
        "signup_end_at": "2026-12-31 23:59", "status": "draft"})
    check("创建计算机批次", ex1.get("code") == 0, ex1.get("message"))
    eid1 = (ex1.get("data") or {}).get("id")

    st, ex2 = call("POST", "/api/exams", {
        "exam_type": "mandarin", "exam_year": 2026, "exam_month": 9,
        "name": "E2E-普通话-%s" % tag, "signup_start_at": "2026-09-01 00:00",
        "signup_end_at": "2026-12-31 23:59", "status": "draft"})
    check("创建普通话批次", ex2.get("code") == 0, ex2.get("message"))
    eid2 = (ex2.get("data") or {}).get("id")

    # 3. 下载导入模板
    print("\n[3] 下载导入模板（openpyxl 写路径）")
    st, blob1 = call("GET", "/api/applications/import-template?exam_id=%d" % eid1, raw=True)
    check("计算机模板可下载", st == 200 and blob1[:2] == b"PK", "status=%s len=%d" % (st, len(blob1)))
    st, blob2 = call("GET", "/api/applications/import-template?exam_id=%d" % eid2, raw=True)
    check("普通话模板可下载", st == 200 and blob2[:2] == b"PK", "status=%s len=%d" % (st, len(blob2)))

    wb1 = load_workbook(io.BytesIO(blob1))
    headers1 = [c.value for c in wb1.worksheets[0][1]]
    check("计算机模板表头 13 列", len(headers1) == 13, len(headers1))
    wb2 = load_workbook(io.BytesIO(blob2))
    headers2 = [c.value for c in wb2.worksheets[0][1]]
    check("普通话模板表头 20 列", len(headers2) == 20, len(headers2))

    # 4. 构造导入文件并上传
    print("\n[4] 批量导入（multipart + openpyxl 读路径）")

    def make_book(headers, rows):
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # 计算机 13 列：机构编码/考点编码/姓名/性别/证件类型/证件号/科目/院校/班级/学历/手机/邮箱/地址
    c_rows = [
        ["3201", "320101", "E2E甲", "男", "1", "320102200301010018",
         "信息处理工程师技术水平", "计算机学院", "计算机2101", "本科",
         "13800001231", "e2e1@test.local", "江苏省南京市鼓楼区"],
        ["3201", "320101", "E2E乙", "女", "1", "320102200301010026",
         "信息处理工程师技术水平", "计算机学院", "计算机2102", "本科",
         "13800001232", "e2e2@test.local", "江苏省南京市玄武区"],
        ["3201", "320101", "E2E丙", "男", "1", "320102200301010034",
         "信息处理工程师技术水平", "计算机学院", "计算机2101", "大专",
         "13800001233", "e2e3@test.local", "江苏省南京市秦淮区"],
    ]
    st, imp1 = upload("/api/applications/import",
                      {"exam_id": str(eid1), "audit_status": "approved", "create_accounts": "1"},
                      "e2e_computer.xlsx", make_book(headers1, c_rows))
    d1 = imp1.get("data") or {}
    check("计算机导入接口返回成功", imp1.get("code") == 0, imp1.get("message"))
    check("计算机导入 3 条全部成功", d1.get("success") == 3,
          "success=%s failed=%s" % (d1.get("success"), d1.get("failed")))
    for e in (d1.get("errors") or [])[:10]:
        print("       失败行 %s: %s" % (e.get("row_no"), e.get("reason")))
    check("计算机导入自动建号 3 个", len(d1.get("created_accounts") or []) == 3,
          len(d1.get("created_accounts") or []))

    # 普通话 20 列
    m_rows = [
        ["E2E丁", "男", "汉族", "1", "320102200301010114", "大学生-语言类", "某大学",
         "13800001241", "2021001", "计算机2101", "计算机学院", "江苏省南京市鼓楼区",
         "江苏省南京市鼓楼区", "210000", "江苏省", "南京市", "鼓楼区",
         "江苏省", "南京市", "鼓楼区"],
        ["E2E戊", "女", "汉族", "1", "320102200301010122", "教师-其他类", "某中学",
         "13800001242", "2021002", "计算机2102", "计算机学院", "江苏省南京市玄武区",
         "江苏省南京市玄武区", "210000", "江苏省", "南京市", "玄武区",
         "江苏省", "南京市", "玄武区"],
    ]
    st, imp2 = upload("/api/applications/import",
                      {"exam_id": str(eid2), "audit_status": "approved", "create_accounts": "1"},
                      "e2e_mandarin.xlsx", make_book(headers2, m_rows))
    d2 = imp2.get("data") or {}
    check("普通话导入接口返回成功", imp2.get("code") == 0, imp2.get("message"))
    check("普通话导入 2 条全部成功", d2.get("success") == 2,
          "success=%s failed=%s" % (d2.get("success"), d2.get("failed")))
    for e in (d2.get("errors") or [])[:10]:
        print("       失败行 %s: %s" % (e.get("row_no"), e.get("reason")))

    # 错误明细（xlsx）下载能力
    st, errbook = call("GET", "/api/applications/import-batches/%d/errors" % d1.get("batch_id"),
                       raw=True)
    check("错误明细可下载为 xlsx", st == 200 and errbook[:2] == b"PK",
          "status=%s len=%d" % (st, len(errbook)))

    # 5. 报名列表
    print("\n[5] 报名列表与导入批次")
    st, lst = call("GET", "/api/applications?exam_id=%d&page=1&page_size=50" % eid1)
    ldata = lst.get("data") or {}
    check("计算机报名列表可访问", lst.get("code") == 0, lst.get("message"))
    check("计算机报名 3 条", ldata.get("total") == 3, "total=%s" % ldata.get("total"))

    st, batches = call("GET", "/api/applications/import-batches?page=1&page_size=5")
    check("导入批次可查询", batches.get("code") == 0, batches.get("message"))
    check("存在导入批次记录", len(((batches.get("data") or {}).get("list")) or []) >= 2)

    # 6. 汇总导出
    print("\n[6] 汇总导出（openpyxl 写路径）")
    st, xb = call("GET", "/api/export/download?exam_id=%d&audit_status=approved" % eid1, raw=True)
    check("计算机导出可下载", st == 200 and xb[:2] == b"PK", "status=%s len=%d" % (st, len(xb)))
    if st == 200 and xb[:2] == b"PK":
        wbx = load_workbook(io.BytesIO(xb))
        wsx = wbx.worksheets[0]
        rowsx = list(wsx.iter_rows(values_only=True))
        check("导出含表头行", len(rowsx) >= 1)
        check("导出数据行 3 条", len(rowsx) - 1 == 3, "rows=%d" % (len(rowsx) - 1))
        check("导出附带辅助 Sheet", len(wbx.sheetnames) >= 2, wbx.sheetnames)

    st, xb2 = call("GET", "/api/export/download?exam_id=%d&audit_status=approved" % eid2, raw=True)
    check("普通话导出可下载", st == 200 and xb2[:2] == b"PK", "status=%s" % st)

    # 7. 清理临时数据
    print("\n[7] 清理本轮临时数据")
    for label, eid in (("计算机", eid1), ("普通话", eid2)):
        st, r = call("DELETE", "/api/exams/%d" % eid)
        check("删除%s临时批次" % label, r.get("code") == 0, r.get("message"))
    st, rs = call("POST", "/api/system/reset",
                  {"scopes": ["applications"], "confirm": "确认清空"})
    check("清理临时报名数据", rs.get("code") == 0, rs.get("message"))

    print("\n" + "=" * 66)
    print("结果：%d 项通过，%d 项失败" % (PASS, FAIL))
    print("=" * 66)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
