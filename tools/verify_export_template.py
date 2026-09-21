# -*- coding: utf-8 -*-
"""需求6：汇总导出「复制官方模板再填充」的列对齐专项验证。

为什么单独写一套
----------------
导出最致命的缺陷不是「导不出来」，而是**导出来了但列错位** —— 表头看着对，
数据整列平移，用的人要等到上级退回才发现。踩坑记录里的真实案例：
同一份官方模板内部 **数据列顺序 ≠ 标题列顺序**（辽宁铁岭/朝阳颠倒、
陕西杨凌示范区排到末列），按列顺序写必错。

之前的覆盖全是间接的：`verify_packaged.py` 只断言模板文件存在、
`compare_samples.py` 是手工工具（还依赖桌面上的文件，跑不进自动化）。
**没有任何一条自动化断言验证「值落在正确的表头下面」** —— 本脚本补的就是这个。

用法：
    python tools/verify_export_template.py http://127.0.0.1:8924

注意：会建批次与报名数据，需**全新库**（数据不落盘，导出的是内存生成的 xlsx）。
"""
import io
import json
import random
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timedelta

from openpyxl import load_workbook

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8924"
PASS, FAIL = [], []

SUFFIX = str(random.randint(10000, 99999))
NAME = f"列校验{SUFFIX}"          # 姓名：唯一的锚点
SCHOOL = f"列校验大学{SUFFIX[:3]}"
CLS = f"校验{SUFFIX[:2]}01"
IDN = "320102199501011239"        # 校验位合法（末位 9）


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method)
    if body is not None:
        req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            blob = r.read()
            return r.status, r.headers, (blob if raw else json.loads(blob.decode("utf-8")))
    except urllib.error.HTTPError as e:
        blob = e.read()
        if raw:
            return e.code, e.headers, blob
        try:
            return e.code, e.headers, json.loads(blob.decode("utf-8"))
        except Exception:
            return e.code, e.headers, {"message": blob[:200].decode("utf-8", "ignore")}


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + ("" if cond else f"  -> {extra}"))


def login(account, password):
    st, _h, r = call("POST", "/api/auth/login", body={"account": account, "password": password})
    return (r["data"]["access_token"] if st == 200 and r.get("data") else None), r


def header_map(ws, max_scan=8):
    """在前 max_scan 行里找出表头行，返回 ({表头: 列号}, 表头行号)。

    表头在官方模板里不一定在第 1 行（可能上面有标题行），必须扫。
    """
    best, best_row, best_hit = {}, 0, 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), 1):
        cells = [("" if v is None else str(v).strip()) for v in row]
        hit = sum(1 for c in cells if c)
        if hit > best_hit:
            best_hit, best_row, best = hit, i, cells
    m = {}
    for idx, c in enumerate(best):
        if c and c not in m:
            m[c] = idx
    return m, best_row


def find_row(ws, keyword, from_row):
    """返回第一条包含 keyword 的行（0 基元组）。"""
    for row in ws.iter_rows(min_row=from_row, values_only=True):
        if any(keyword in ("" if v is None else str(v)) for v in row):
            return row
    return None


def cell_of(row, hmap, *names):
    """按表头名取该行的值；表头名可能有多种写法，逐个试。"""
    for n in names:
        for k, idx in hmap.items():
            if n in k or k in n:
                if idx < len(row):
                    v = row[idx]
                    return "" if v is None else str(v).strip()
    return None


def window(days=30):
    """报名窗口必须覆盖「当前时间」，否则提交会被「不在报名开放期」拦下。"""
    now = datetime.now()
    return ((now - timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S"),
            (now + timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S"))


def main():
    print("=" * 70)
    print(f"导出模板列对齐专项验证 · {BASE}\n")

    print("[0] 准备")
    admin, r = login("admin", "Admin@123")
    check("管理员登录成功", bool(admin), r)
    if not admin:
        return 1

    st, _h, r = call("GET", "/api/dicts/org", token=admin)
    orgs = r.get("data") or []
    check("取得考试机构字典", bool(orgs), r)
    st, _h, r = call("GET", "/api/dicts/subject", token=admin)
    subjects = r.get("data") or []
    check("取得报考科目字典", bool(subjects), r)

    # 建批次（open，允许报名）
    # 批次名由「年/月」生成，演示数据里通常已有当前月 → 用两年后的随机月份避免重名
    s0, s1 = window()
    y, m = datetime.now().year + 2, random.randint(1, 12)
    st, _h, r = call("POST", "/api/exams", token=admin, body={
        "exam_type": "computer", "exam_year": y, "exam_month": m,
        "signup_start_at": s0, "signup_end_at": s1,
        "description": "列对齐校验批次", "status": "open"})
    check("建计算机批次成功", st == 200 and r.get("code") == 0, r)
    if st != 200:
        return 1
    exam_id = r["data"]["id"]

    # 建考生并提交报名
    uname = f"tpl{SUFFIX}"
    st, _h, r = call("POST", "/api/users", token=admin, body={
        "username": uname, "phone": f"137{random.randint(10000000, 99999999)}",
        "role": "candidate", "password": "Tpl@12345"})
    check("建考生账号成功", st == 200 and r.get("code") == 0, r)
    cand, r = login(uname, "Tpl@12345")
    check("考生登录成功", bool(cand), r)

    st, _h, r = call("POST", "/api/applications", token=cand, body={
        "exam_id": exam_id, "org_code": orgs[0]["code"], "exam_site_code": "110101",
        "name": NAME, "gender": "男", "id_type": "1", "id_number": IDN,
        "subject": subjects[0]["name"] if isinstance(subjects[0], dict) else subjects[0],
        "school": SCHOOL, "class_name": CLS, "education": "本科",
        "phone": "13800002222", "email": "tpl@exam.local",
        "address": "江苏省南京市鼓楼区1号"})
    check("报名提交成功", st == 200 and r.get("code") == 0, r)
    if st != 200:
        return 1
    app_id = r["data"]["id"]

    # app_type 必填：审核路由要据此选表（计算机/普通话字段集不同）
    st, _h, r = call("PUT", f"/api/applications/{app_id}/audit", token=admin,
                     body={"app_type": "computer", "action": "approved", "comment": ""})
    check("审核通过（导出默认只取已通过）", st == 200 and r.get("code") == 0, r)

    print("\n[1] 模板模式：值必须落在正确的表头下面")
    q = urllib.parse.urlencode({"exam_id": exam_id, "audit_status": "approved",
                                "table_mode": "template"})
    st, hdr, blob = call("GET", "/api/export/download?" + q, token=admin, raw=True)
    check("模板模式导出成功", st == 200 and blob[:2] == b"PK", (st, blob[:40]))
    if st != 200 or blob[:2] != b"PK":
        return 1

    wb = load_workbook(io.BytesIO(blob))
    ws = wb[wb.sheetnames[0]]
    hmap, hrow = header_map(ws)
    check("能定位到表头行", bool(hmap), f"表头 {hmap}")
    check("表头含「姓名」列", any("姓名" in k for k in hmap), str(list(hmap))[:200])

    row = find_row(ws, NAME, hrow + 1)
    check("导出内容含本次报名", row is not None, f"未找到 {NAME}")
    if row is None:
        return 1

    # 以下每一条都在防「整列平移」：值必须出现在**同名表头**那一列
    checks = [
        ("姓名", NAME, ("姓名",)),
        ("证件号码", IDN, ("证件号码", "身份证号", "证件号")),
        # 计算机模板里这一列叫「就读或者毕业院校」，普通话模板另有叫法 → 都列上
        ("学校 / 院系", SCHOOL, ("就读或者毕业院校", "学校", "院系")),
        ("班级", CLS, ("班级",)),
        ("性别", "男", ("性别",)),
    ]
    for label, expect, names in checks:
        got = cell_of(row, hmap, *names)
        check(f"「{label}」的值落在同名表头列下", got == expect, f"期望 {expect!r} 实得 {got!r}")

    # 反向验证：把数据行整体右移一列，上面的断言必须失配。
    # 少了这一步，「按列顺序写」这种错位 bug 有可能因为取值逻辑写得宽松而蒙混过关。
    misaligned = (None,) + tuple(row[:-1])
    check("反向验证：数据错位一列后断言会失败（证明断言有鉴别力）",
          cell_of(misaligned, hmap, "姓名") != NAME,
          f"错位后仍取到 {cell_of(misaligned, hmap, '姓名')!r}")

    print("\n[2] 自建模式（需求6 保留的备选）")
    q2 = urllib.parse.urlencode({"exam_id": exam_id, "audit_status": "approved",
                                 "table_mode": "build"})
    st2, _h, blob2 = call("GET", "/api/export/download?" + q2, token=admin, raw=True)
    check("自建模式导出成功", st2 == 200 and blob2[:2] == b"PK", (st2, blob2[:40]))
    if st2 == 200 and blob2[:2] == b"PK":
        wb2 = load_workbook(io.BytesIO(blob2))
        ws2 = wb2[wb2.sheetnames[0]]
        hmap2, hrow2 = header_map(ws2)
        row2 = find_row(ws2, NAME, hrow2 + 1)
        check("自建模式同样含该条数据", row2 is not None, f"未找到 {NAME}")
        if row2 is not None:
            got = cell_of(row2, hmap2, "姓名")
            check("自建模式姓名列对齐", got == NAME, f"期望 {NAME!r} 实得 {got!r}")
        # 两种模式产出不同（说明确实走了两条路径，不是同一个结果）
        check("两种模式产出不同内容", blob != blob2, "模板模式与自建模式字节相同")

    # 注：通用类型没有官方模板 → 自动退回自建，这条由 verify_generic.py §4 覆盖
    # （它导出的是带自定义字段的通用批次，能成功即证明回退没炸）。

    print("\n" + "=" * 70)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    if FAIL:
        for f in FAIL:
            print("  FAILED:", f)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
