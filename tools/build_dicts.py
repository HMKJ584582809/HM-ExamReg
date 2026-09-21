# -*- coding: utf-8 -*-
"""
从两份官方报名模板解析字典与表头，生成 backend/app/data/dicts.json。

- 报名导入模板-计算机.xlsx : Sheet1 表头(13列) / 填表说明 / sheet3 机构字典+科目字典
- 考生报名模板-普通话.xlsx : Sheet1 表头(20列) / Sheet2 职业字典 / Sheet3 省市区三级 / Sheet4 模板版本

用法：
    python tools/build_dicts.py <官方模板所在目录>
    EXAM_TEMPLATE_DIR=<目录> python tools/build_dicts.py

考试机构只保留编码、名称统一写成「示例机构-<编码>」：dicts.json 会打进对外分发的 exe，
真实机构名没必要跟着扩散。
"""
import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# 默认目录（可被命令行参数 / EXAM_TEMPLATE_DIR 覆盖）。
# 别写死到某个人的桌面：一是换机器就跑不了，二是路径里带着具体目录名也没必要进仓库。
SRC = Path(__file__).resolve().parents[1] / "backend" / "app" / "assets" / "templates"
DST = Path(__file__).resolve().parents[1] / "backend" / "app" / "data" / "dicts.json"

# 证件类型（两张模板说明一致）
ID_TYPES = [
    {"code": "1", "name": "居民身份证"},
    {"code": "6", "name": "香港居民身份证"},
    {"code": "7", "name": "澳门居民身份证"},
    {"code": "8", "name": "台湾居民身份证"},
]

ETHNICITIES = [
    "汉族", "蒙古族", "回族", "藏族", "维吾尔族", "苗族", "彝族", "壮族", "布依族", "朝鲜族",
    "满族", "侗族", "瑶族", "白族", "土家族", "哈尼族", "哈萨克族", "傣族", "黎族", "傈僳族",
    "佤族", "畲族", "高山族", "拉祜族", "水族", "东乡族", "纳西族", "景颇族", "柯尔克孜族",
    "土族", "达斡尔族", "仫佬族", "羌族", "布朗族", "撒拉族", "毛南族", "仡佬族", "锡伯族",
    "阿昌族", "普米族", "塔吉克族", "怒族", "乌孜别克族", "俄罗斯族", "鄂温克族", "德昂族",
    "保安族", "裕固族", "京族", "塔塔尔族", "独龙族", "鄂伦春族", "赫哲族", "门巴族",
    "珞巴族", "基诺族", "其他",
]

EDUCATIONS = ["博士研究生", "硕士研究生", "本科", "专科", "中专", "高中", "初中", "其他"]


def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return "" if v is None else str(v).strip()


def parse_computer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    headers = [cell(ws, 1, c) for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h]

    ws_note = wb["填表说明"]
    instructions = []
    for r in range(1, ws_note.max_row + 1):
        v = ws_note.cell(row=r, column=1).value
        if v is None:
            continue
        text = str(v).replace("\r\n", "\n").strip()
        if text:
            instructions.append(text)

    ws3 = wb["sheet3"]
    orgs, subjects = [], []
    # 官方 sheet3：第 1 行留空、第 2 行是表头（机构编码/考试机构名称/报考科目），
    # 数据从第 3 行开始——从第 2 行读会把表头当成一条机构记录。
    for r in range(3, ws3.max_row + 1):
        code = cell(ws3, r, 2)
        name = cell(ws3, r, 3)
        subj = cell(ws3, r, 5)
        if code and name:
            # ⚠ 机构名称不落真实名：dicts.json 会被打进 exe 对外分发，
            #   真实机构名属于第三方机构信息，没必要随成品扩散。
            #   保留编码（唯一标识、业务校验依赖它），名称换成明显占位。
            orgs.append({"code": code, "name": f"示例机构-{code}"})
        if subj:
            subjects.append(subj)

    # 去重保序
    subjects = list(dict.fromkeys(subjects))
    return {
        "headers": headers,
        "instructions": instructions,
        "orgs": orgs,
        "subjects": subjects,
    }


def parse_mandarin(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    headers = [cell(ws, 1, c) for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h]

    ws2 = wb["Sheet2"]
    occupations = []
    for r in range(1, ws2.max_row + 1):
        v = cell(ws2, r, 1)
        if v:
            occupations.append(v)

    ws4 = wb["Sheet4"]
    version = cell(ws4, 2, 1)

    ws3 = wb["Sheet3"]
    max_row, max_col = ws3.max_row, ws3.max_column

    # 省级字典：A 列（跳过表头与注释行）
    provinces_flat = []
    for r in range(3, max_row + 1):
        v = cell(ws3, r, 1)
        if v:
            provinces_flat.append(v)

    # 定位每个省的列块：第 1 行 "市级（省名）"
    blocks = []
    for c in range(1, max_col + 1):
        h = cell(ws3, 1, c)
        if h.startswith("市级（") and h.endswith("）"):
            blocks.append((c, h[3:-1]))
    blocks.append((max_col + 1, None))  # 哨兵

    regions = []
    used_cols = set()      # 已被省份块占用的列（含市级列与各城市数据列）
    for i in range(len(blocks) - 1):
        start, prov = blocks[i]
        end = blocks[i + 1][0]
        cities = []
        for r in range(3, max_row + 1):
            v = cell(ws3, r, start)
            if v:
                cities.append(v)
        city_nodes = [{"name": c, "counties": []} for c in cities]
        name_to_node = {c["name"]: c for c in city_nodes}
        used_cols.add(start)
        # 块内其后每一列对应一个城市：第 3 行是城市名，第 4 行起是县区。
        # 按【名称】匹配，不要按顺序匹配——官方模板个别省份的数据列顺序与市级列
        # 并不一致（如辽宁省「铁岭市 / 朝阳市」顺序颠倒），按顺序匹配会在中途中断，
        # 导致该省后半段城市整段丢失数据。
        for col in range(start + 1, end):
            city_name = cell(ws3, 3, col)
            node = name_to_node.get(city_name) if city_name else None
            if node is None:
                continue
            node["counties"] = [cell(ws3, r, col) for r in range(4, max_row + 1)
                                if cell(ws3, r, col)]
            node["src_col"] = col
            used_cols.add(col)
        regions.append({"province": prov, "cities": city_nodes})

    # 补全「省份块内没有数据列」的条目：官方模板把个别省直辖单位单独排在最后一列
    # （如陕西省「杨凌示范区」位于末列）。从未被任何省份块占用的列中按第 3 行名称
    # 反查并标记 standalone，导出时单独占一列，避免该省块多占列导致后续省份错位。
    for region in regions:
        for node in region["cities"]:
            if node["counties"]:
                continue
            for col in range(1, max_col + 1):
                if col in used_cols or cell(ws3, 3, col) != node["name"]:
                    continue
                counties = [cell(ws3, r, col) for r in range(4, max_row + 1)
                            if cell(ws3, r, col)]
                if counties:
                    node["counties"] = counties
                    node["standalone"] = True
                    node["src_col"] = col
                    used_cols.add(col)
                    break

    return {
        "headers": headers,
        "occupations": occupations,
        "version": version,
        "provinces": provinces_flat,
        "regions": regions,
    }


def main():
    # ⚠ 源码里不该写死某个人的桌面路径（那也是隐私，而且换台机器就跑不了）。
    #   优先用命令行参数 / 环境变量，找不到再退回仓库内的 assets/templates。
    # ⚠ Path("") 是 "."、真值，会短路掉后面的默认值——必须先判空字符串
    env = os.environ.get("EXAM_TEMPLATE_DIR", "").strip()
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else (Path(env) if env else SRC)

    def pick(*names):
        """按候选名序找第一个存在的文件（官方模板原名 / 仓库内简名）。"""
        for n in names:
            p = src / n
            if p.exists():
                return p
        raise SystemExit(f"找不到模板文件，请在 {src} 下放 {names[0]}")

    computer = parse_computer(pick("报名导入模板-计算机.xlsx", "computer.xlsx"))
    mandarin = parse_mandarin(pick("考生报名模板-普通话.xlsx", "mandarin.xlsx"))

    data = {
        "id_types": ID_TYPES,
        "ethnicities": ETHNICITIES,
        "educations": EDUCATIONS,
        "computer": computer,
        "mandarin": mandarin,
    }

    DST.parent.mkdir(parents=True, exist_ok=True)
    DST.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    print("dicts.json ->", DST)
    print("计算机表头", len(computer["headers"]), computer["headers"])
    print("计算机机构", len(computer["orgs"]), "科目", len(computer["subjects"]))
    print("普通话表头", len(mandarin["headers"]), mandarin["headers"])
    print("普通话职业", len(mandarin["occupations"]), "版本", mandarin["version"])
    print("省级", len(mandarin["provinces"]), "省级块", len(mandarin["regions"]))
    total_city = sum(len(r["cities"]) for r in mandarin["regions"])
    total_county = sum(len(c["counties"]) for r in mandarin["regions"] for c in r["cities"])
    print("市级合计", total_city, "县区合计", total_county)
    for r in mandarin["regions"][:4]:
        print("  ", r["province"], len(r["cities"]), "市",
              sum(len(c["counties"]) for c in r["cities"]), "县区",
              "| 首市:", r["cities"][0]["name"] if r["cities"] else None,
              len(r["cities"][0]["counties"]) if r["cities"] else 0)


if __name__ == "__main__":
    main()
