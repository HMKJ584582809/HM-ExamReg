# -*- coding: utf-8 -*-
"""解析两份官方报名模板：输出 sheet 结构、表头、字典数据。"""
import json
import sys
from pathlib import Path

import openpyxl

BASE = Path(r"C:\Users\Administrator\Desktop\考试报名采集核对")
OUT = Path(r"C:\Users\Administrator\WorkBuddy AI\2026-09-16-16-50-46\build\template_dump")
OUT.mkdir(parents=True, exist_ok=True)


def dump(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    result = {"file": path.name, "sheets": []}
    for ws in wb.worksheets:
        sheet = {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merged": [str(r) for r in ws.merged_cells.ranges][:40],
            "rows": [],
        }
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > 60:
                break
            vals = ["" if v is None else str(v).strip() for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if not vals:
                continue
            sheet["rows"].append({"r": i, "v": vals})
        result["sheets"].append(sheet)
    return result


if __name__ == "__main__":
    all_data = []
    for f in ["报名导入模板-计算机.xlsx", "考生报名模板-普通话.xlsx"]:
        p = BASE / f
        d = dump(p)
        all_data.append(d)
        print("=" * 100)
        print("FILE:", f)
        for s in d["sheets"]:
            print("-" * 80)
            print("SHEET:", s["title"], "rows=", s["max_row"], "cols=", s["max_col"])
            print("MERGED:", s["merged"])
            for r in s["rows"]:
                print(r["r"], "|", " | ".join(r["v"]))
    (OUT / "dump.json").write_text(json.dumps(all_data, ensure_ascii=False, indent=2), encoding="utf-8")
