# -*- coding: utf-8 -*-
"""比对「导出示例」与官方模板的结构一致性（Sheet 名 / 表头 / 辅助 Sheet 非空单元格）"""
import sys
from pathlib import Path

from openpyxl import load_workbook

DESK = Path(r"C:\Users\Administrator\Desktop\考试报名采集核对")

PAIRS = [
    ("普通话", DESK / "考生报名模板-普通话.xlsx",
     DESK / "发布" / "导出示例-普通话水平测试-考生报名模板.xlsx"),
    ("计算机", DESK / "报名导入模板-计算机.xlsx",
     DESK / "发布" / "导出示例-计算机类考试-报名导入模板.xlsx"),
]


def grid(ws):
    return [tuple("" if v is None else str(v).strip() for v in row)
            for row in ws.iter_rows(values_only=True)]


def nonempty(ws):
    out = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), 1):
        for c, v in enumerate(row, 1):
            if v is not None and str(v).strip() != "":
                out[(r, c)] = str(v).strip()
    return out


ok = True
for label, official, exported in PAIRS:
    print("=" * 70)
    print("%s  ·  官方模板 vs 导出示例" % label)
    print("  官方：%s" % official.name)
    print("  导出：%s" % exported.name)
    if not official.exists() or not exported.exists():
        print("  [FAIL] 文件缺失")
        ok = False
        continue

    wo, we = load_workbook(official), load_workbook(exported)
    same_sheets = wo.sheetnames == we.sheetnames
    print("  Sheet 列表：%s" % ("一致" if same_sheets else "不一致"))
    print("     官方 = %s" % wo.sheetnames)
    print("     导出 = %s" % we.sheetnames)
    ok = ok and same_sheets

    for name in wo.sheetnames:
        if name not in we.sheetnames:
            continue
        so, se = wo[name], we[name]
        go, ge = grid(so), grid(se)
        hdr_same = bool(go) and bool(ge) and go[0] == ge[0]
        print("  [%s] 表头一致：%s" % (name, "是" if hdr_same else "否"))
        if not hdr_same:
            print("       官方 = %s" % (go[0] if go else None,))
            print("       导出 = %s" % (ge[0] if ge else None,))
            ok = False
        # 辅助 Sheet 逐格比对
        if name != wo.sheetnames[0]:
            no, ne = nonempty(so), nonempty(se)
            missing = {k: v for k, v in no.items() if k not in ne}
            extra = {k: v for k, v in ne.items() if k not in no}
            diff = {k: (no[k], ne[k]) for k in no if k in ne and no[k] != ne[k]}
            print("       非空单元格：官方 %d / 导出 %d；缺失 %d，多余 %d，值不同 %d"
                  % (len(no), len(ne), len(missing), len(extra), len(diff)))
            for k, v in list(missing.items())[:3]:
                print("         缺失 %s: %r" % (k, v))
            for k, v in list(diff.items())[:3]:
                print("         不同 %s: 官方=%r 导出=%r" % (k, v[0], v[1]))
            if missing or extra or diff:
                ok = False

print("=" * 70)
print("结论：" + ("导出示例与官方模板结构完全一致 ✅" if ok else "存在差异 ❌"))
sys.exit(0 if ok else 1)
