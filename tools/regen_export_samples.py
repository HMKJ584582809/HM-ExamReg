# -*- coding: utf-8 -*-
"""
离线重新生成导出示例文件（不依赖运行中的服务与演示数据）

用途：字典或导出布局调整后，用官方模板结构重新生成 `导出示例-*.xlsx`。
会保留原示例 Sheet1 中的数据行（若文件已存在），仅重建表头样式与辅助 Sheet。

用法：
  python tools/regen_export_samples.py <输出目录>
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import load_workbook                       # noqa: E402

from app.routers.applications import (                   # noqa: E402
    COMPUTER_FIELDS, MANDARIN_FIELDS)
from app.routers.export import build_workbook            # noqa: E402

SPECS = [
    ("computer", "导出示例-计算机类考试-报名导入模板.xlsx", COMPUTER_FIELDS),
    ("mandarin", "导出示例-普通话水平测试-考生报名模板.xlsx", MANDARIN_FIELDS),
]


def read_rows(src: Path, fields: list):
    """从既有示例文件中读回 Sheet1 数据行（跳过表头）。"""
    if not src.exists():
        return []
    ws = load_workbook(src).worksheets[0]
    it = ws.iter_rows(values_only=True)
    next(it, None)
    rows = []
    for row in it:
        item = {}
        for i, f in enumerate(fields):
            v = row[i] if i < len(row) else None
            item[f] = "" if v is None else str(v)
        if any(item.values()):
            rows.append(item)
    return rows


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    out.mkdir(parents=True, exist_ok=True)
    for et, fname, fields in SPECS:
        target = out / fname
        rows = read_rows(target, fields)
        exam = {"id": 0, "exam_type": et, "name": target.stem}
        build_workbook(exam, rows, target, len(rows))
        print("已生成：%s（数据行 %d）" % (target.name, len(rows)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
