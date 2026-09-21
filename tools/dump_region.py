# -*- coding: utf-8 -*-
"""逐列解析普通话模板 Sheet3，弄清省市区三级排布规则。"""
from pathlib import Path

import openpyxl

p = Path(r"C:\Users\Administrator\Desktop\考试报名采集核对\考生报名模板-普通话.xlsx")
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb["Sheet3"]
print("dims", ws.max_row, ws.max_column)
print("merged sample:", [str(r) for r in list(ws.merged_cells.ranges)[:20]])
print("merged count:", len(ws.merged_cells.ranges))

# 找出所有合并区域，判断列分组
merged_by_col = {}
for rng in ws.merged_cells.ranges:
    merged_by_col.setdefault(rng.min_col, []).append(str(rng))
for k in sorted(merged_by_col):
    print("col", k, merged_by_col[k][:5])

print("=" * 90)
for col in range(1, 40):
    vals = []
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=col).value
        if v is not None and str(v).strip():
            vals.append((row, str(v).strip()))
    if vals:
        print(f"--- COL {col} ({openpyxl.utils.get_column_letter(col)}) n={len(vals)}")
        print("   ", vals[:30])
